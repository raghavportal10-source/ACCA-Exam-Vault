"""Convert FM_Mock_Exam_01.md into print-ready, black-and-white Word files.

Usage: python3 md_to_docx.py <exam.md> <questions.docx> <answers.docx> [<combined.docx>]

The markdown is split at the double rule ("---" on two consecutive lines):
everything before it is the question paper, everything after it is the answer
key. Formatting is black and white only, with native Word headings and bold,
and a solid paragraph-border rule after every question block.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
FONT = "Arial"
CONTENT_CM = 21.0 - 4.4


def _set_font(style):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class Writer:
    def __init__(self, footer_text):
        self.doc = Document()
        self.block_open = False
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10.5)
        st.font.color.rgb = BLACK
        _set_font(st)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.12
        for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11), ("List Bullet", 10.5)):
            s = self.doc.styles[name]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.color.rgb = BLACK
            s.font.italic = False
            _set_font(s)
            if name.startswith("Heading"):
                s.font.bold = True
                s.paragraph_format.space_before = Pt(12)
                s.paragraph_format.space_after = Pt(6)
                s.paragraph_format.keep_with_next = True
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for txt, fld in ((f"{footer_text}, page ", None), (None, "PAGE"), (" of ", None), (None, "NUMPAGES")):
            r = fp.add_run(txt or "")
            r.font.size = Pt(8)
            r.font.color.rgb = BLACK
            if fld:
                _field(r, fld)

    # ----- inline text -----
    @staticmethod
    def runs(p, text, size=None, bold=False):
        # **bold**, *italic*, `answer box`
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("*"):
                r = p.add_run(part[1:-1])
                r.italic = True
                r.bold = bold or None
            elif part.startswith("`"):
                inner = part[1:-1]
                for j, piece in enumerate(re.split(r"(_{3,})", inner)):
                    if not piece:
                        continue
                    if piece.startswith("___"):
                        r = p.add_run(" " * max(len(piece) * 3, 24))
                        r.underline = True
                    else:
                        r = p.add_run(piece)
                        r.bold = True
            else:
                r = p.add_run(part)
                r.bold = bold or None
            if size:
                r.font.size = Pt(size)
        return p

    def para(self, text="", size=None, align=None, after=4, before=0, keep=False, indent=None, bold=False):
        p = self.doc.add_paragraph()
        self.runs(p, text, size, bold)
        pf = p.paragraph_format
        pf.space_after = Pt(after)
        pf.space_before = Pt(before)
        pf.keep_with_next = keep
        if indent is not None:
            pf.left_indent = Cm(indent)
        if align is not None:
            p.alignment = align
        return p

    # ----- blocks -----
    def rule(self):
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)
        p.paragraph_format.space_after = Pt(8)

    def close_block(self):
        if self.block_open:
            self.rule()
            self.block_open = False

    def heading(self, text, level, page_break=False):
        self.close_block()
        h = self.doc.add_heading("", level=level)
        self.runs(h, text)
        for r in h.runs:
            r.bold = True
            r.font.color.rgb = BLACK
        h.paragraph_format.page_break_before = page_break
        return h

    def question(self, num, marks=2):
        self.close_block()
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"Question {num}")
        r.bold = True
        r = p.add_run(f"\t({marks} marks)")
        r.font.size = Pt(9)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), alignment=2)
        self.block_open = True

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        self.runs(p, text)
        p.paragraph_format.space_after = Pt(2)

    def table(self, rows):
        header, body = rows[0], rows[2:]
        rows = [header] + body
        ncol = len(header)
        clean = lambda s: re.sub(r"[*`]", "", s)
        lens = [max(len(clean(r[i])) if i < len(r) else 0 for r in rows) for i in range(ncol)]
        weights = [min(max(l, 5), 60) for l in lens]
        widths = [CONTENT_CM * w / sum(weights) for w in weights]
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for ri, row in enumerate(rows):
            trPr = t.rows[ri]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            for ci in range(ncol):
                val = row[ci].strip() if ci < len(row) else ""
                c = t.cell(ri, ci)
                c.width = Cm(widths[ci])
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                self.runs(p, val, size=12 if val == "☐" else 9.5, bold=(ri == 0))
                numeric = re.fullmatch(r"[\d$(),.%\s]+(days|units)?", clean(val) or "x") is not None
                if val == "☐" or (ci > 0 and (numeric or ri == 0)):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ci, w in enumerate(widths):
            t.columns[ci].width = Cm(w)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path, title):
        self.close_block()
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        self.doc.core_properties.title = title
        self.doc.core_properties.author = "ACCA Exam Vault"
        self.doc.save(path)


def render(w, lines, part):
    """Render markdown lines. part is 'questions' or 'answers'."""
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "---":
            w.close_block()
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip().strip("|").split("|"))
                i += 1
            w.table(rows)
            continue
        m = re.match(r"^(#{1,3}) (.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1 and part == "questions" and first_heading:
                w.para("", after=50)
                w.para("ACCA", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
                w.para("**Financial Management (FM)**", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
                w.para("Mock Examination 01", size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
            elif level == 1:
                w.heading(text, 1, page_break=not first_heading or part == "answers")
            elif level == 2:
                w.heading(text, 1, page_break=text.startswith(("Section A", "Section B")))
            else:
                if part == "answers" and re.match(r"Q\d+:", text):
                    w.heading("Question " + text[1:], 3)
                    w.block_open = True
                else:
                    w.heading(text, 2)
            first_heading = False
            i += 1
            continue
        m = re.match(r"^\*\*(\d+)\.\*\* (.*)$", s)
        if m:
            w.question(m.group(1))
            w.para(m.group(2), after=5, keep=True)
            i += 1
            continue
        m = re.match(r"^- (.*)$", s)
        if m:
            text = m.group(1)
            if re.match(r"^[A-D]\. ", text):
                w.para(f"**{text[0]}**\t{text[3:]}", indent=0.6, after=2, keep=True)
            elif text.startswith("☐"):
                p = w.para("", indent=0.6, after=2, keep=True)
                r = p.add_run("☐")
                r.font.size = Pt(12)
                w.runs(p, "   " + text[1:].strip())
            else:
                w.bullet(text)
            i += 1
            continue
        m = re.match(r"^(\d+)\. (.*)$", s)
        if m:
            p = w.para(f"({m.group(1)})\t{m.group(2)}", indent=1.2, after=2, keep=True)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            i += 1
            continue
        w.para(s, after=5)
        i += 1


def main():
    src, out_q, out_a = sys.argv[1:4]
    out_c = sys.argv[4] if len(sys.argv) > 4 else None
    lines = open(src, encoding="utf-8").read().split("\n")
    split = next(k for k in range(len(lines) - 1) if lines[k].strip() == "---" and lines[k + 1].strip() == "---")
    q_lines, a_lines = lines[:split], lines[split + 2:]

    wq = Writer("ACCA FM Mock Exam 01: Question Paper")
    render(wq, q_lines, "questions")
    wq.para("**End of question paper**", align=WD_ALIGN_PARAGRAPH.CENTER, before=12)
    wq.save(out_q, "ACCA FM Mock Exam 01: Question Paper")

    wa = Writer("ACCA FM Mock Exam 01: Solutions")
    wa.para("", after=50)
    wa.para("ACCA", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    wa.para("**Financial Management (FM)**", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    wa.para("Mock Examination 01: Solutions and Examiner Rationale", size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    wa.para("For use with the question paper FM_Mock_Exam_01_Questions.docx", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    render(wa, a_lines, "answers")
    wa.save(out_a, "ACCA FM Mock Exam 01: Solutions")

    if out_c:
        wc = Writer("ACCA FM Mock Exam 01")
        render(wc, q_lines, "questions")
        wc.close_block()
        render(wc, a_lines, "answers")
        wc.save(out_c, "ACCA FM Mock Exam 01")


if __name__ == "__main__":
    main()
