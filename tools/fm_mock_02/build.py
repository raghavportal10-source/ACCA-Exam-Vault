"""Build ACCA FM Mock Exam 02 (100 marks) as two print-ready Word files with python-docx.

Usage: python3 build.py <questions.docx> <answers.docx>
(The earlier combined FM_Mock_Exam_02_100Marks.docx is no longer generated.)

Formatting is strictly black and white: no colour, no shading, and a solid
black rule (a native Word paragraph border) after every question block.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from q31 import df, q31

OUT_Q, OUT_A = sys.argv[1], sys.argv[2]
BLACK = RGBColor(0, 0, 0)
ACCENT = GREY = BLACK
FONT = "Arial"
CONTENT_CM = 21.0 - 4.4

doc = None
_block_open = False


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _set_font(style):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def start_doc(footer_text):
    """Create a fresh A4 document with black-and-white styles and a page-numbered footer."""
    global doc, _block_open
    doc = Document()
    _block_open = False
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = BLACK
    _set_font(st)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.12
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11), ("List Bullet", 10.5)):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.color.rgb = BLACK
        s.font.italic = False
        _set_font(s)
        if name.startswith("Heading"):
            s.font.bold = True
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for txt, fld in ((f"{footer_text}, page ", None), (None, "PAGE"), (" of ", None), (None, "NUMPAGES")):
        r = fp.add_run(txt or "")
        r.font.size = Pt(8)
        r.font.color.rgb = BLACK
        if fld:
            _field(r, fld)


def finish(path, title):
    close_block()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = title
    doc.core_properties.author = "ACCA Exam Vault"
    doc.save(path)


def rule():
    """Solid black horizontal line: a native bottom border on an empty paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)


def open_block():
    global _block_open
    _block_open = True


def close_block():
    global _block_open
    if _block_open:
        rule()
        _block_open = False


# ---------- helpers ----------
def add_runs(p, text, size=None, color=None, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**")
        r = p.add_run(part[2:-2] if bold else part)
        r.bold = bold or None
        r.italic = italic or None
        if size:
            r.font.size = Pt(size)
    return p


def P(text="", size=None, align=None, after=4, before=0, keep=False, indent=None, italic=False, color=None):
    p = doc.add_paragraph()
    add_runs(p, text, size, None, italic)
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.keep_with_next = keep
    if indent is not None:
        pf.left_indent = Cm(indent)
    if align:
        p.alignment = align
    return p


def H(text, level=1, page_break=False):
    close_block()
    h = doc.add_heading(text, level=level)
    h.paragraph_format.page_break_before = page_break
    for r in h.runs:
        r.font.color.rgb = BLACK
    return h


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(2)
    if level:
        p.paragraph_format.left_indent = Cm(1.27 + 0.63 * level)
    return p


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def T(rows, widths=None, header=True, size=9.5, center=None, bold_rows=(), right=None, after=8):
    """rows: list of lists of str. widths in cm (summing to <= content width). Header row is bold, unshaded."""
    ncol = len(rows[0])
    if widths is None:
        first = CONTENT_CM * (0.46 if ncol > 2 else 0.6)
        widths = [first] + [(CONTENT_CM - first) / (ncol - 1)] * (ncol - 1)
    center = set(center if center is not None else range(1, ncol))
    right = set(right or [])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        no_split(t.rows[ri])
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.width = Cm(widths[ci])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, str(val), size=size if val != "☐" else 12)
            if (header and ri == 0) or ri in bold_rows:
                for r in p.runs:
                    r.bold = True
            if ci in right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif ci in center or val == "☐":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ci, w in enumerate(widths):
        t.columns[ci].width = Cm(w)
    if header:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(after)
    return t


def Q(num, marks=2):
    """Start a question block: number on the left, marks on the right."""
    close_block()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Question {num}")
    r.bold = True
    r = p.add_run(f"\t({marks} marks)")
    r.font.size = Pt(9)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), alignment=2)  # right
    open_block()
    return p


def stem(text, keep=True):
    return P(text, keep=keep, after=5)


def opts(items):
    for i, it in enumerate(items):
        P(f"**{'ABCD'[i]}**\t{it}", indent=0.6, after=2, keep=i < len(items) - 1)
    P("", after=2)


def checks(items):
    for i, it in enumerate(items):
        p = P("", indent=0.6, after=2, keep=i < len(items) - 1)
        r = p.add_run("☐")
        r.font.size = Pt(12)
        add_runs(p, "   " + it)
    P("", after=2)


def statements(items):
    for i, it in enumerate(items, 1):
        p = P(f"({i})\t{it}", indent=1.2, after=2, keep=True)
        p.paragraph_format.first_line_indent = Cm(-0.8)


def answer_box(prefix="", suffix=""):
    p = P("", indent=0.6, after=8)
    r = p.add_run(f"{prefix} ")
    r.bold = True
    r2 = p.add_run(" " * 40)
    r2.underline = True
    if suffix:
        p.add_run(f"  {suffix}").bold = True


def tf(rows, cols=("True", "False")):
    T([["Statement", *cols]] + [[s] + ["☐"] * len(cols) for s in rows],
      widths=[CONTENT_CM - 2.3 * len(cols)] + [2.3] * len(cols))


def money(x, dp=0):
    v = round(x, dp)
    s = f"{abs(v):,.{dp}f}"
    return f"({s})" if v < 0 else s


# =====================================================================
# QUESTION PAPER
# =====================================================================
start_doc("ACCA FM Mock Exam 02: Question Paper")
P("", after=60)
P("ACCA", size=14, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("**Financial Management (FM)**", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
P("Mock Examination 02: 100 Marks", size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, color=ACCENT)
P("Coverage: Study Text Chapters 1 to 10", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=30, color=GREY)
P("**Time allowed: 3 hours**", align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
T([["Section", "Content", "Marks"],
   ["A", "15 objective test questions, 2 marks each. ALL are compulsory.", "30"],
   ["B", "3 scenarios with 5 objective test questions each, 2 marks each. ALL are compulsory.", "30"],
   ["C", "2 constructed-response questions, 20 marks each. BOTH are compulsory.", "40"],
   ["**Total**", "", "**100**"]],
  widths=[2.2, CONTENT_CM - 4.6, 2.4], center={0, 2})
P("Instructions to candidates", size=11, after=4, before=12, color=ACCENT).runs[0].bold = True
for line in (
    "Answer every question. There is no negative marking.",
    "Objective test questions are worth 2 marks each and are marked right or wrong. Multiple-response, True/False and matching questions must be answered in full to earn the marks.",
    "Where a numeric answer is required, follow the rounding instruction given in the question.",
    "Use the formulae and the present value and annuity tables provided at the end of the question paper (3 decimal places).",
    "Assume a 365-day year unless the question states otherwise.",
    "In Section C, show all workings. Marks are awarded for method as well as for the final answer.",
):
    bullet(line)
page_break()

# =====================================================================
# SECTION A
# =====================================================================
H("Section A: ALL 15 questions are compulsory and MUST be attempted", 1)
P("Each question is worth 2 marks.", italic=True, color=GREY)

Q("1")
stem("Which stakeholder group is primarily concerned with a company's ability to pay interest on time and to repay capital when it falls due?")
opts(["Ordinary shareholders", "Long-term lenders", "Employees", "Customers"])

Q("2")
stem("Identify whether each of the following statements about government economic policy is true or false.")
tf(["An increase in interest rates by a central bank will usually cause the country's currency to appreciate",
    "Expansionary fiscal policy involves increasing taxation to reduce aggregate demand",
    "Monetary policy is concerned with the money supply, interest rates and the availability of credit"])

Q("3")
stem("Halloran Treasury Co has bought a 91-day Treasury bill for $98.40 per $100 nominal value and will hold it to maturity.")
stem("What is the effective annual yield on the Treasury bill, assuming a 365-day year and that the return can be reinvested on the same terms (to **two decimal places**)?")
answer_box("", "%")

Q("4")
stem("The following information relates to Vantage Moorings Co for the year just ended:")
T([["", "$m"], ["Profit before interest and tax", "4.20"], ["Interest payable", "0.60"],
   ["Preference dividends paid", "0.30"]], widths=[9.0, 3.0])
stem("Corporation tax is 25% of profit before tax. There are 12 million ordinary shares in issue and the current share price is $3.15.")
stem("What is Vantage Moorings Co's price/earnings ratio?")
opts(["10.5 times", "12.0 times", "14.0 times", "15.75 times"])

Q("5")
stem("Which TWO of the following actions would shorten a company's cash operating cycle?")
checks(["Reducing the average time that raw materials are held in inventory",
        "Paying suppliers within 10 days to take an early settlement discount, instead of the normal 45 days",
        "Negotiating 60-day credit terms with suppliers instead of 30 days",
        "Increasing the credit period offered to customers to boost sales"])

Q("6")
stem("Galloway Sprockets Co uses a reorder level system to control its inventory of component G4. Data for component G4:")
T([["", "Minimum", "Average", "Maximum"],
   ["Usage (units per week)", "400", "600", "900"],
   ["Supply lead time (weeks)", "2", "3", "4"]])
stem("The reorder quantity is the economic order quantity of 5,000 units. The reorder level is set so that there is no risk of a stock-out.")
stem("What is the maximum inventory level for component G4 (in **units**)?")
answer_box("", "units")

Q("7")
stem("Strathmore Joinery Co is considering paying its suppliers well beyond their agreed credit terms in order to ease pressure on its overdraft.")
stem("Which TWO of the following are risks to Strathmore Joinery Co of taking this action?")
checks(["Damage to the company's credit rating with credit reference agencies",
        "A longer cash operating cycle",
        "Suppliers may refuse to supply on credit in the future or may demand payment on delivery",
        "An increase in the company's inventory holding period"])

Q("8")
stem("Pelham Interiors Co needs $2.4 million of cash a year for its operations, spread evenly over the year. It meets this need by selling short-term securities which earn interest at 4% per year. Each sale of securities costs $60.")
stem("Using the Baumol model, what is the optimal amount of securities to sell each time cash is needed (to the **nearest $**)?")
answer_box("$")

Q("9")
stem("Which of the following is **NOT** a disadvantage of the payback period as an investment appraisal method?")
opts(["It ignores the time value of money",
      "It ignores cash flows that arise after the payback period",
      "It is simple to calculate and easy for managers to understand",
      "It does not measure the effect of a project on shareholder wealth"])

Q("10")
stem("Ravensholt Estates Co expects to receive $45,000 per year in perpetuity from a property investment, with the first receipt at the end of Year 4. The company's cost of capital is 9% per year.")
stem("What is the present value of the receipts?")
opts(["$354,000", "$386,000", "$421,000", "$500,000"])

Q("11")
stem("Kittering Coaches Co is considering an investment of $250,000 in new equipment. The equipment will generate net cash inflows of $68,000 per year for five years and will be sold for $20,000 at the end of Year 5. The cost of capital is 11% per year.")
stem("What is the net present value of the investment (to the **nearest $**)?")
answer_box("$")

Q("12")
stem("Sefton Robotics Co will buy a machine for $120,000 at the start of Year 1. Tax-allowable depreciation is available at 25% per year on a reducing balance basis, with the first allowance claimed in Year 1. Corporation tax is 30%, paid one year in arrears, and the after-tax cost of capital is 10%.")
stem("What is the present value of the tax benefit arising from the tax-allowable depreciation claimed in **Year 2**?")
opts(["$5,069", "$5,576", "$6,750", "$6,759"])

Q("13")
stem("Identify whether each of the following statements about inflation in investment appraisal is true or false.")
tf(["If all cash flows inflate at the general rate of inflation, discounting real cash flows at the real cost of capital gives the same NPV as discounting money cash flows at the money cost of capital",
    "Tax-allowable depreciation should be increased each year by the general rate of inflation",
    "Inflation will increase the working capital requirement of a project even if its sales volumes stay constant"])

Q("14")
stem("Marston Leisure Co has estimated the following possible net present values for a new venue:")
T([["Outcome", "NPV ($)", "Probability"], ["Strong demand", "120,000", "0.40"],
   ["Moderate demand", "30,000", "0.35"], ["Weak demand", "(80,000)", "0.25"]])
stem("What is the expected net present value of the venue (to the **nearest $**)?")
answer_box("$")

Q("15")
stem("Branscombe Water Co has $500,000 of capital available for investment and cannot raise more. The following projects are **indivisible** and cannot be delayed. Projects A and D are mutually exclusive.")
T([["Project", "Initial investment ($'000)", "NPV ($'000)"], ["A", "200", "70"], ["B", "250", "80"],
   ["C", "150", "48"], ["D", "300", "105"]])
stem("What is the maximum NPV that Branscombe Water Co can generate?")
opts(["$128,000", "$150,000", "$153,000", "$175,000"])

# =====================================================================
# SECTION B
# =====================================================================
H("Section B: ALL 15 questions are compulsory and MUST be attempted", 1, page_break=True)
P("Each question is worth 2 marks.", italic=True, color=GREY)

H("The following scenario relates to questions 16 to 20", 2)
P("Halberd Textiles Co sells workwear to retailers. All of its annual sales of $21.9 million are on credit, and customers currently take an average of 40 days to pay. Bad debts are 1% of sales. The company's contribution margin is 25% of sales.")
P("The sales director has proposed extending the credit period offered to all customers to 60 days. This is expected to increase sales by 10%. All customers, existing and new, are expected to take the full 60 days. Bad debts are expected to rise to 1.5% of total sales. There will be no change in fixed costs or inventory levels.")
P("Halberd Textiles Co finances its working capital with an overdraft at an interest rate of 8% per year.")
P("Extracts from its latest statement of financial position are:")
T([["", "$m"], ["Inventory", "1.8"], ["Trade receivables", "2.4"], ["Cash", "0.3"],
   ["Trade payables", "1.9"], ["Bank overdraft", "1.1"]], widths=[9.0, 3.0])

Q("16")
stem("What is the annual increase in the cost of financing trade receivables if the proposal is accepted?")
opts(["$19,200", "$96,000", "$124,800", "$316,800"])

Q("17")
stem("What is the net annual financial effect of the proposal on Halberd Textiles Co's profit?")
opts(["Increase of $280,350", "Increase of $389,850", "Increase of $405,150", "Increase of $422,700"])

Q("18")
stem("What is Halberd Textiles Co's current quick (acid test) ratio (to **two decimal places**)?")
answer_box("", ": 1")

Q("19")
stem("Identify whether each of the following statements about working capital investment policies is true or false.")
tf(["A conservative policy involves holding high levels of inventory and cash and offering generous credit to customers",
    "An aggressive policy is likely to give a higher return on capital employed, but with a greater risk of liquidity problems",
    "The appropriate level of working capital investment is unaffected by the industry in which a company operates"])

Q("20")
stem("Halberd Textiles Co prepares an aged analysis of trade receivables each month. What is the main purpose of this analysis?")
opts(["To calculate the company's trade receivables collection period",
      "To identify overdue debts so that collection can be prioritised",
      "To assess the creditworthiness of new customers before credit is granted",
      "To calculate the allowance for bad debts required by accounting standards"])

H("The following scenario relates to questions 21 to 25", 2)
P("Kestrel Ridge Co is evaluating a four-year project to manufacture a new range of industrial filters. The project requires a machine costing $750,000 at the start of Year 1. The machine is specialised and is expected to be sold for $350,000 at the end of Year 4.")
P("Forecast sales volumes are:")
T([["Year", "1", "2", "3", "4"], ["Sales volume (units)", "40,000", "45,000", "50,000", "35,000"]])
for line in (
    "The selling price in current (Year 0) terms is $30 per unit, which will increase by 4% per year from Year 1.",
    "The variable cost in current terms is $18 per unit, which will increase by 5% per year from Year 1.",
    "Working capital equal to 15% of each year's sales revenue is required at the **start** of that year and is recovered in full at the end of the project.",
    "Tax-allowable depreciation is available at 25% per year on a reducing balance basis in Years 1 to 3. No 25% allowance is claimed in Year 4, the year of disposal; instead a balancing allowance or charge equal to the difference between the tax written-down value at the start of Year 4 and the disposal proceeds arises in Year 4.",
    "Corporation tax is 30% per year, paid one year in arrears.",
    "Kestrel Ridge Co's real cost of capital is 6.8% per year and general inflation is expected to be 3% per year.",
):
    bullet(line)

Q("21")
stem("What is Kestrel Ridge Co's money (nominal) cost of capital (to **one decimal place**)?")
answer_box("", "%")

Q("22")
stem("What is the project's contribution in Year 2, in money terms (to the **nearest $**)?")
opts(["$540,000", "$553,500", "$567,135", "$572,886"])

Q("23")
stem("What is the working capital cash flow in **Year 1** of the project?")
opts(["$(219,024)", "$(187,200)", "$(31,824)", "$31,824"])

Q("24")
stem("What is the tax effect of disposing of the machine, and in which year does it arise?")
opts(["A balancing allowance giving a tax saving of $10,078 in Year 5",
      "A balancing charge giving additional tax of $10,078 in Year 5",
      "A balancing charge giving additional tax of $10,078 in Year 4",
      "A balancing charge giving additional tax of $105,000 in Year 5"])

Q("25")
stem("Which of the following statements about Kestrel Ridge Co's appraisal are correct?")
statements(["If general inflation turns out higher than expected but the project's money cash flows stay as forecast, the NPV will fall",
            "Fixed costs should always be inflated at the general rate of inflation, even where a specific rate is available",
            "Discounting money cash flows at the real cost of capital would overstate the NPV"])
opts(["1 and 2 only", "1 and 3 only", "2 and 3 only", "1, 2 and 3"])

H("The following scenario relates to questions 26 to 30", 2)
P("Montclair Freight Co is a listed logistics company. Its board is reviewing how the stock market values the company and is comparing its performance with the sector. Extracts from the latest financial statements are:")
T([["", "$m"], ["Profit before interest and tax", "14.4"], ["Interest", "(2.4)"], ["Profit before tax", "12.0"],
   ["Tax at 20%", "(2.4)"], ["Profit after tax", "9.6"], ["Preference dividends", "(0.6)"],
   ["Ordinary dividends", "(5.4)"], ["Retained profit for the year", "3.6"]], widths=[9.0, 3.0], bold_rows=(5, 8))
for line in (
    "There are 30 million ordinary shares in issue, with a current market price of $4.50 per share.",
    "The book value of ordinary shareholders' equity is $80 million.",
    "Montclair Freight Co has $30 million (nominal value) of 8% loan notes in issue, currently trading at $108 per $100 nominal value.",
    "The sector average price/earnings ratio is 14 times. Analysts expect Montclair Freight Co's earnings attributable to ordinary shareholders to grow by 8% next year.",
):
    bullet(line)

Q("26")
stem("What is Montclair Freight Co's current dividend yield (to **one decimal place**)?")
answer_box("", "%")

Q("27")
stem("If the market applied the sector average price/earnings ratio to Montclair Freight Co's **forecast** earnings per share for next year, what would its share price be?")
opts(["$4.20", "$4.54", "$4.56", "$4.84"])

Q("28")
stem("Montclair Freight Co's finance director is studying the yield curve before refinancing the loan notes. Identify whether each of the following statements is true or false.")
tf(["Liquidity preference theory helps to explain why a normal yield curve slopes upwards",
    "An inverted yield curve suggests that the market expects interest rates to rise in the future",
    "Market segmentation theory suggests that the short and long ends of the yield curve are determined by different groups of investors"])

Q("29")
stem("Montclair Freight Co raises part of its finance through banks and other financial intermediaries. Which TWO of the following are functions of financial intermediaries?")
checks(["Maturity transformation, by lending long-term from short-term deposits",
        "Guaranteeing a minimum return to the companies that borrow from them",
        "Risk reduction, by pooling deposits and lending to many borrowers",
        "Removing the need for companies to publish financial statements"])

Q("30")
stem("What are Montclair Freight Co's gearing ratio (debt/equity, based on **market values**) and interest cover?")
T([["Option", "Gearing (debt/equity, market values)", "Interest cover"],
   ["A", "24.0%", "6.0 times"], ["B", "37.5%", "6.0 times"], ["C", "24.0%", "5.0 times"], ["D", "19.4%", "6.0 times"]],
  widths=[2.0, 8.0, CONTENT_CM - 10.0])

# =====================================================================
# SECTION C
# =====================================================================
H("Section C: BOTH questions are compulsory and MUST be attempted", 1, page_break=True)

COST, SCRAP = 2_900_000, 300_000
npv, rows, cf, pv, twdv3 = q31(COST, SCRAP)
npv_ovh = q31(COST, SCRAP, ovh=90_000)[0]

H("Question 31: Brindlewood Aggregates Co", 2)
P("Brindlewood Aggregates Co produces construction materials. It is considering a four-year project to make a new low-carbon paving block. The finance director has prepared a draft appraisal, which included the costs of market research already carried out and a share of head office overheads, and which showed a negative NPV. You have been asked to prepare a revised appraisal.")
for line in (
    "The project requires new machinery costing $2,900,000 at the start of Year 1. The machinery will be sold for $300,000 at the end of Year 4.",
    "Market research costing $75,000 was completed and paid for last month.",
    "The selling price in current (Year 0) terms is $48 per block, which will increase by 3% per year from Year 1.",
    "The variable cost in current terms is $27 per block, which will increase by 5% per year from Year 1.",
    "Incremental fixed production costs are $520,000 per year in current terms, increasing by 4% per year from Year 1.",
    "The project will be charged a share of existing head office overheads of $90,000 per year in current terms. These overheads will be incurred whether or not the project goes ahead.",
    "Working capital equal to 10% of each year's sales revenue is required at the **start** of that year and is recovered in full at the end of Year 4.",
    "Tax-allowable depreciation is available at 25% per year on a reducing balance basis in Years 1 to 3. In Year 4, the year of disposal, no 25% allowance is claimed; instead a balancing allowance or charge arises equal to the tax written-down value at the start of Year 4 less the disposal proceeds.",
    "Corporation tax is 25% per year, paid one year in arrears.",
    "Brindlewood Aggregates Co uses a nominal after-tax cost of capital of 11% per year to appraise projects.",
):
    bullet(line)
P("Forecast sales volumes are:", before=4)
T([["Year", "1", "2", "3", "4"], ["Sales volume (blocks)", "60,000", "85,000", "95,000", "70,000"]])
P("**Required:**", keep=True)
for lab, txt, mk in (
    ("(a)", "Calculate the net present value of the project. Work to the nearest $.", "14 marks"),
    ("(b)", "Advise the board whether the project should be undertaken, commenting on the treatment of the market research and head office overheads in the finance director's draft appraisal. Discuss how sensitivity analysis and probability analysis could help the board to assess the risk of the project.", "6 marks"),
):
    p = P(f"**{lab}**\t{txt}", indent=1.0, after=4)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), alignment=2)
    r = p.add_run(f"\t({mk})")
    r.bold = True
P("**(20 marks)**", align=WD_ALIGN_PARAGRAPH.RIGHT)
open_block()

H("Question 32: Calloway Hardware Co", 2, page_break=True)
P("Calloway Hardware Co is a wholesaler of tools and building supplies. All of its sales and purchases are on credit. It has grown quickly and its bank overdraft, which finances all of its current assets, has reached $5.8 million. The bank has expressed concern about the company's liquidity.")
P("Extracts from the latest financial statements:")
T([["", "$m"], ["Revenue", "36.50"], ["Cost of sales (equal to credit purchases)", "25.55"], ["Inventory", "4.20"],
   ["Trade receivables", "6.00"], ["Trade payables", "3.50"], ["Bank overdraft", "5.80"]], widths=[9.0, 3.0])
P("Sector averages: cash operating cycle 45 days; current ratio 1.50 times; quick ratio 0.90 times.")
P("A factoring company has offered to take over the administration of Calloway Hardware Co's sales ledger on a **non-recourse** basis, on the following terms:")
for line in (
    "The factor will charge an annual fee of 1.5% of revenue.",
    "The factor expects to reduce the average trade receivables collection period to 40 days.",
    "The factor will advance 80% of trade receivables at an interest rate of 10% per year. The remaining receivables will continue to be financed by the overdraft.",
    "Calloway Hardware Co's bad debts, currently 1% of revenue, will be eliminated.",
    "Credit control administration costs of $180,000 per year will be saved.",
    "The overdraft interest rate is 9% per year. Assume a 365-day year.",
):
    bullet(line)
P("**Required:**", keep=True)
for lab, txt, mk in (
    ("(a)", "Calculate Calloway Hardware Co's cash operating cycle, current ratio and quick ratio, and compare them briefly with the sector averages.", "4 marks"),
    ("(b)", "Calculate whether the factoring offer is financially acceptable to Calloway Hardware Co.", "8 marks"),
    ("(c)", "Explain the difference between aggressive, conservative and matching policies for financing working capital, and identify which policy Calloway Hardware Co is currently following.", "4 marks"),
    ("(d)", "Discuss the non-financial factors that Calloway Hardware Co should consider before accepting the factoring offer.", "4 marks"),
):
    p = P(f"**{lab}**\t{txt}", indent=1.0, after=4)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), alignment=2)
    r = p.add_run(f"\t({mk})")
    r.bold = True
P("**(20 marks)**", align=WD_ALIGN_PARAGRAPH.RIGHT)
open_block()

# =====================================================================
# FORMULAE AND TABLES
# =====================================================================
H("Formulae sheet", 1, page_break=True)
for f in (
    "**Economic order quantity** = √(2C₀D / C_H)",
    "**Miller-Orr model:** Return point = Lower limit + (1/3 × spread); Spread = 3 × [(3/4 × transaction cost × variance of cash flows) / interest rate]^(1/3)",
    "**Baumol model:** Optimal amount of cash to raise = √(2 × transaction cost × annual cash requirement / interest rate)",
    "**Fisher formula:** (1 + i) = (1 + r)(1 + h)",
    "**Present value of a perpetuity** of $1 a year from Year 1 = 1 / r",
    "**Equivalent annual cost** = PV of costs / annuity factor for the life of the cycle",
    "**Annual cost of an early settlement discount** = (100 / (100 − d))^(365 / t) − 1",
):
    bullet(f)

def tables(kind):
    for lo, hi in ((1, 10), (11, 20)):
        rates = list(range(lo, hi + 1))
        rows = [["Year"] + [f"{r}%" for r in rates]]
        for n in range(1, 11):
            if kind == "pv":
                rows.append([str(n)] + [f"{(1 + r / 100) ** -n:.3f}" for r in rates])
            else:
                rows.append([str(n)] + [f"{sum((1 + r / 100) ** -k for k in range(1, n + 1)):.3f}" for r in rates])
        w = (CONTENT_CM - 1.2) / 10
        T(rows, widths=[1.2] + [w] * 10, size=7.5, center=range(0, 11), after=6)

H("Present value table", 2)
P("Present value of $1 = (1 + r)^(−n), where r = discount rate and n = number of periods until payment.", size=9, color=GREY)
tables("pv")
H("Annuity table", 2, page_break=True)
P("Present value of an annuity of $1 per year = [1 − (1 + r)^(−n)] / r", size=9, color=GREY)
tables("af")
P("**End of question paper**", align=WD_ALIGN_PARAGRAPH.CENTER, before=12)

# =====================================================================
# ANSWER KEY
# =====================================================================
finish(OUT_Q, "ACCA FM Mock Exam 02: Question Paper")

start_doc("ACCA FM Mock Exam 02: Solutions")
P("", after=60)
P("ACCA", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("**Financial Management (FM)**", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
P("Mock Examination 02: Solutions and Marking Guide", size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
P("For use with the question paper FM_Mock_Exam_02_Questions.docx", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
P("**Contents**", after=4)
for line in ("Answer summary for Sections A and B", "Section A: worked solutions and distractor rationale (Questions 1 to 15)",
             "Section B: worked solutions and distractor rationale (Questions 16 to 30)",
             "Section C: full calculations, marking guides and model answers (Questions 31 and 32)"):
    bullet(line)
H("Solution Key and Examiner Rationale", 1, page_break=True)
P("Difficulty mix by marks: Hard 50, Average 30, Easy 20. In Sections A and B (60 marks), 14 questions are Hard, 10 Average and 6 Easy. In Section C (40 marks), 22 marks are Hard, 10 Average and 8 Easy.", italic=True, color=GREY)

H("Answer summary: Sections A and B", 2)
summary = [
    ("1", "B", "Easy", "MCQ", "1 FM function"),
    ("2", "True, False, True", "Average", "True/False", "2 FM environment"),
    ("3", "6.68%", "Hard", "Numeric", "2 FM environment"),
    ("4", "D", "Hard", "MCQ", "1 FM function"),
    ("5", "1st and 3rd", "Average", "Multi-response", "3 Working capital"),
    ("6", "7,800 units", "Hard", "Numeric", "4 Managing WC"),
    ("7", "1st and 3rd", "Average", "Multi-response", "4 Managing WC"),
    ("8", "$84,853", "Average", "Numeric", "5 WC finance"),
    ("9", "C", "Easy", "MCQ", "6 Investment decisions"),
    ("10", "B", "Hard", "MCQ", "7 DCF"),
    ("11", "$13,188", "Average", "Numeric", "7 DCF"),
    ("12", "A", "Hard", "MCQ", "8 Tax and inflation"),
    ("13", "True, False, True", "Hard", "True/False", "8 Tax and inflation"),
    ("14", "$38,500", "Average", "Numeric", "9 Risk"),
    ("15", "C", "Hard", "MCQ", "10 Specific decisions"),
    ("16", "C", "Hard", "MCQ", "4 Managing WC"),
    ("17", "A", "Hard", "MCQ", "4 Managing WC"),
    ("18", "0.90", "Easy", "Numeric", "3 Working capital"),
    ("19", "True, True, False", "Average", "True/False", "3 Working capital"),
    ("20", "B", "Easy", "MCQ", "4 Managing WC"),
    ("21", "10.0%", "Easy", "Numeric", "8 Tax and inflation"),
    ("22", "C", "Average", "MCQ", "8 Tax and inflation"),
    ("23", "C", "Hard", "MCQ", "8 Tax and inflation"),
    ("24", "B", "Hard", "MCQ", "8 Tax and inflation"),
    ("25", "B", "Average", "Statement combination", "8 Tax and inflation"),
    ("26", "4.0%", "Easy", "Numeric", "1 FM function"),
    ("27", "B", "Hard", "MCQ", "1 FM function"),
    ("28", "True, False, True", "Hard", "True/False", "2 FM environment"),
    ("29", "1st and 3rd", "Average", "Multi-response", "2 FM environment"),
    ("30", "A", "Hard", "Grid MCQ", "1 FM function"),
]
T([["Q", "Answer", "Level", "Format", "Chapter"]] + [list(s) for s in summary],
  widths=[1.0, 3.8, 2.0, 4.0, CONTENT_CM - 10.8], size=8.5, center={0, 2})

H("Section A", 2, page_break=True)

def key(title, lines):
    H(title, 3)
    open_block()
    for ln in lines:
        if isinstance(ln, list):
            T(ln[1:] if ln[0] == "table" else ln, size=9)
        elif ln.startswith("- "):
            bullet(ln[2:])
        else:
            P(ln)

key("Question 1: B (Long-term lenders)", [
    "Lenders are concerned with the company's ability to service and repay debt. That is why they look at interest cover and gearing, and why loan agreements include covenants.",
    "- **A:** shareholders are mainly concerned with dividends and share price growth, not repayment of capital.",
    "- **C and D:** employees care about pay and job security, and customers about quality, price and continuity of supply.",
])
key("Question 2: True, False, True", [
    "- **True.** Higher interest rates attract inflows of foreign capital, which increases demand for the currency.",
    "- **False.** This is the trap. Expansionary fiscal policy *reduces* taxation and/or increases government spending to *raise* aggregate demand. Increasing taxation is contractionary.",
    "- **True.** This is the definition of monetary policy.",
])
key("Question 3: 6.68%", [
    "Return for 91 days = (100 − 98.40) / 98.40 = 1.626%",
    "Effective annual yield = (100 / 98.40)^(365/91) − 1 = **6.68%**",
    "Distractors:",
    "- 6.52%: simple annualisation, 1.626% × 365/91, which ignores compounding.",
    "- 6.42%: discount yield, (1.60 / 100) × 365/91, which uses the nominal value instead of the price paid.",
])
key("Question 4: D (15.75 times)", [
    "Earnings attributable to ordinary shareholders = (4.20 − 0.60) × 75% − 0.30 = 2.70 − 0.30 = $2.40m",
    "EPS = 2.40 / 12 = 20.0 cents, so P/E = 3.15 / 0.20 = **15.75 times**",
    "Distractors:",
    "- **C (14.0):** forgets to deduct preference dividends (EPS 22.5c).",
    "- **B (12.0):** taxes PBIT without deducting interest (EPS 26.25c).",
    "- **A (10.5):** uses profit before tax (EPS 30c).",
])
key("Question 5: 1st and 3rd", [
    "Cash operating cycle = inventory days + receivables days − payables days.",
    "- **Less time in inventory: correct.** It shortens the cycle.",
    "- **Longer supplier credit: correct.** More payables days reduce the cycle.",
    "- **Paying suppliers in 10 days: wrong.** It *reduces* payables days, so the cycle gets longer. The discount may still be worth taking, but that is a separate cost/benefit decision.",
    "- **Longer customer credit: wrong.** It increases receivables days.",
])
key("Question 6: 7,800 units", [
    "Reorder level = maximum usage × maximum lead time = 900 × 4 = 3,600 units",
    "Maximum inventory = reorder level + reorder quantity − (minimum usage × minimum lead time) = 3,600 + 5,000 − (400 × 2) = **7,800 units**",
    "Traps:",
    "- 8,600: leaves out the deduction for minimum usage during the minimum lead time.",
    "- 6,800: deducts average usage × average lead time (600 × 3) instead of the minimums.",
])
key("Question 7: 1st and 3rd", [
    "- **Credit rating damage: correct.** Persistent late payment is reported to credit agencies and damages the company's rating.",
    "- **Suppliers refusing credit: correct.** Suppliers may withdraw credit, insist on cash on delivery, or give the company lower priority.",
    "- **Longer cash operating cycle: wrong.** Paying later *shortens* the cycle. This is the benefit the company is seeking, not a risk.",
    "- **Longer inventory holding period: wrong.** Payment terms do not directly affect inventory days.",
])
key("Question 8: $84,853", [
    "Q = √(2 × 60 × 2,400,000 / 0.04) = √7,200,000,000 = **$84,853**",
    "At this amount, annual transaction costs (2,400,000 / 84,853 × $60 = $1,697) equal the interest forgone on the average cash balance (84,853 / 2 × 4% = $1,697).",
    "Traps:",
    "- Entering the interest rate as 4 instead of 0.04 gives $8,485.",
    "- Omitting the 2 in the formula gives $60,000.",
])
key("Question 9: C", [
    "Simplicity is an *advantage* of payback, so it is the only option that is not a disadvantage.",
    "- **A, B and D** are all standard criticisms of payback.",
])
key("Question 10: B ($386,000)", [
    "The perpetuity formula values a stream starting one year later, so a stream starting in Year 4 is valued at Year 3: 45,000 / 0.09 = $500,000",
    "Discount back three years: PV = 500,000 × 0.772 = **$386,000**",
    "An alternative method gives the same answer to rounding: 500,000 − (45,000 × annuity factor for Years 1 to 3 of 2.531) = $386,105.",
    "Distractors:",
    "- **A ($354,000):** discounts with the Year 4 factor (0.708). This is the classic off-by-one trap.",
    "- **C ($421,000):** discounts with the Year 2 factor.",
    "- **D ($500,000):** no discounting at all.",
])
key("Question 11: $13,188", [
    "NPV = −250,000 + 68,000 × 3.696 + 20,000 × 0.593 = −250,000 + 251,328 + 11,860 = **$13,188**",
    "Trap: leaving out the scrap value gives $1,328.",
])
key("Question 12: A ($5,069)", [
    "Year 2 TAD = 120,000 × 75% × 25% = $22,500",
    "Tax saving = 22,500 × 30% = $6,750, received in **Year 3** because tax is paid one year in arrears.",
    "PV = 6,750 × 0.751 = **$5,069**",
    "Distractors:",
    "- **B ($5,576):** discounts with the Year 2 factor, ignoring the arrears timing.",
    "- **C ($6,750):** not discounted.",
    "- **D ($6,759):** uses the Year 1 allowance (30,000 × 30% × 0.751).",
])
key("Question 13: True, False, True", [
    "- **True.** The real and money methods give the same NPV when every cash flow inflates at the general rate.",
    "- **False.** Tax-allowable depreciation is based on the historic cost of the asset and is **not** inflated. This is a common error in Section C.",
    "- **True.** Working capital is held in money terms, so rising prices increase the amount needed even if volumes are unchanged. The increase is an extra cash outflow each year.",
])
key("Question 14: $38,500", [
    "ENPV = (0.40 × 120,000) + (0.35 × 30,000) + (0.25 × −80,000) = 48,000 + 10,500 − 20,000 = **$38,500**",
    "Note: the ENPV is not a possible outcome. There is also a 25% chance of a negative NPV, which the ENPV hides.",
])
key("Question 15: C ($153,000)", [
    "Indivisible projects mean the feasible combinations within $500,000 must be compared:",
    ["table", ["Combination", "Cost ($'000)", "NPV ($'000)"], ["A + B", "450", "150"], ["A + C", "350", "118"],
     ["B + C", "400", "128"], ["C + D", "450", "**153**"], ["A + D", "500", "175 (not allowed: mutually exclusive)"]],
    "Distractors:",
    "- **D ($175,000):** ignores the mutual exclusivity. Ranking by profitability index (D and A both 0.35) gives the same wrong combination, and PI ranking is only valid for divisible projects anyway.",
])

H("Section B", 2, page_break=True)
key("Question 16: C ($124,800)", [
    "Current receivables = 21.9m × 40/365 = $2.40m",
    "New sales = 21.9m × 1.10 = $24.09m, so new receivables = 24.09m × 60/365 = $3.96m",
    "Increase in receivables = $1.56m × 8% = **$124,800**",
    "Distractors:",
    "- **B ($96,000):** applies 60 days to the *old* sales figure.",
    "- **D ($316,800):** finances the whole new balance instead of the increase.",
    "- **A ($19,200):** grows sales but keeps 40 days.",
])
key("Question 17: A (Increase of $280,350)", [
    "Extra contribution = 2.19m × 25% = $547,500",
    "Extra bad debts = (1.5% × 24.09m) − (1% × 21.9m) = 361,350 − 219,000 = $(142,350)",
    "Extra financing cost (Q16) = $(124,800)",
    "Net benefit = **$280,350**",
    "Distractors:",
    "- **B:** charges 1.5% bad debts only on the *new* sales. The higher rate applies to all sales.",
    "- **C:** ignores the financing cost.",
    "- **D:** ignores the bad debts.",
])
key("Question 18: 0.90", [
    "Quick ratio = (receivables + cash) / current liabilities = (2.4 + 0.3) / (1.9 + 1.1) = 2.7 / 3.0 = **0.90**",
    "Traps:",
    "- Including inventory gives the *current* ratio, 1.50.",
    "- Leaving the overdraft out of current liabilities gives 1.42.",
])
key("Question 19: True, True, False", [
    "- **True.** A conservative policy holds high levels of inventory and cash and offers generous credit.",
    "- **True.** An aggressive policy keeps the investment lower, which raises ROCE but increases liquidity risk.",
    "- **False.** Industry norms are a key influence. A supermarket and a construction company need very different levels of working capital.",
])
key("Question 20: B", [
    "An aged analysis groups balances by how long they have been outstanding, so overdue accounts can be chased first.",
    "- **A:** the collection period is calculated from receivables and sales, not from an aged analysis.",
    "- **C:** new customers have no balances to analyse.",
    "- **D:** an aged analysis may help *inform* the allowance for bad debts, but that is not its main purpose for credit control.",
])
key("Question 21: 10.0%", [
    "(1 + i) = 1.068 × 1.03 = 1.10004, so **i = 10.0%**",
    "Trap: simple addition (6.8% + 3%) gives 9.8%.",
])
key("Question 22: C ($567,135)", [
    "Year 2 selling price = 30 × 1.04² = $32.448; Year 2 variable cost = 18 × 1.05² = $19.845",
    "Contribution = (32.448 − 19.845) × 45,000 = **$567,135**",
    "Distractors:",
    "- **B ($553,500):** inflates for one year only.",
    "- **D ($572,886):** inflates the current contribution of $12 at the general rate of 3%.",
    "- **A ($540,000):** no inflation.",
])
key("Question 23: C ($(31,824))", [
    "Year 1 sales = 40,000 × 30 × 1.04 = $1,248,000, so working capital needed at Year 0 = $187,200",
    "Year 2 sales = 45,000 × 30 × 1.04² = $1,460,160, so working capital needed at Year 1 = $219,024",
    "Year 1 cash flow = increase of $31,824, an **outflow**",
    "Distractors:",
    "- **A:** the total requirement instead of the increase.",
    "- **B:** the Year 0 flow.",
    "- **D:** wrong sign.",
])
key("Question 24: B (Balancing charge, $10,078 in Year 5)", [
    ["table", ["", "$"], ["Cost", "750,000"], ["Year 1 TAD (25%)", "(187,500)"], ["Year 2 TAD", "(140,625)"],
     ["Year 3 TAD", "(105,469)"], ["TWDV at start of Year 4", "316,406"], ["Disposal proceeds", "(350,000)"],
     ["**Balancing charge (Year 4)**", "**33,594**"]],
    "Additional tax = 33,594 × 30% = **$10,078**, payable in **Year 5** (one year in arrears).",
    "As stated in the scenario, no 25% allowance is claimed in Year 4. The balancing adjustment replaces it, so total net relief over the asset's life is cost less proceeds: $750,000 − $350,000 = $400,000.",
    "Distractors:",
    "- **A:** this is a *charge*, not an allowance, because proceeds exceed the TWDV. Too much relief has been claimed and it is clawed back.",
    "- **C:** wrong year.",
    "- **D:** taxes the whole sale proceeds.",
])
key("Question 25: B (1 and 3 only)", [
    "- **(1) True.** Higher inflation raises the money cost of capital (Fisher), so fixed money cash flows are discounted more heavily.",
    "- **(2) False.** Specific inflation rates should be used where they are available.",
    "- **(3) True.** The real rate is lower than the money rate, so applying it to inflated money flows overstates the NPV.",
])
key("Question 26: 4.0%", [
    "DPS = 5.4m / 30m = 18 cents, so dividend yield = 0.18 / 4.50 = **4.0%**",
])
key("Question 27: B ($4.54)", [
    "Current EPS = (9.6 − 0.6) / 30 = 30.0 cents; forecast EPS = 30.0 × 1.08 = 32.4 cents",
    "Share price = 32.4c × 14 = **$4.54**",
    "Distractors:",
    "- **A ($4.20):** uses current EPS instead of the forecast.",
    "- **C ($4.56):** grows profit *before* preference dividends, then deducts them.",
    "- **D ($4.84):** forgets to deduct preference dividends.",
])
key("Question 28: True, False, True", [
    "- **True.** Investors want extra return for tying up funds for longer, so long-term yields are higher.",
    "- **False.** An inverted curve (long-term yields below short-term yields) suggests rates are expected to **fall**.",
    "- **True.** Under market segmentation theory, different investor groups (for example banks at the short end, pension funds at the long end) set each end of the curve.",
])
key("Question 29: 1st and 3rd", [
    "Financial intermediaries also aggregate small deposits into larger loans.",
    "- **Guaranteeing borrowers a return: wrong.** Intermediaries do not guarantee returns to borrowers.",
    "- **Removing the need to publish accounts: wrong.** Publication of financial statements is a legal and regulatory requirement, which intermediaries do not affect.",
])
key("Question 30: A", [
    "Market value of debt = 30m × 1.08 = $32.4m; market value of equity = 30m × $4.50 = $135m",
    "Gearing (debt/equity) = 32.4 / 135 = **24.0%**; interest cover = PBIT / interest = 14.4 / 2.4 = **6.0 times**",
    "Distractors:",
    "- **B:** uses book values (30 / 80).",
    "- **C:** uses profit before tax for interest cover.",
    "- **D:** uses debt / (debt + equity), which is a different gearing measure from the one asked for.",
])

# ---------- Section C solutions ----------
H("Section C", 2, page_break=True)
H("Question 31: Brindlewood Aggregates Co", 3)
open_block()
P("**(a) Net present value** (all figures in $, money terms)")

yrs = ["0", "1", "2", "3", "4", "5"]
def row(label, vals_by_year):
    return [label] + [money(vals_by_year.get(t, 0)) if t in vals_by_year else "" for t in range(6)]

sales, vcs, fcs, op = rows["sales"], rows["vc"], rows["fc"], rows["op"]
tbl = [["Year"] + yrs,
       row("Sales revenue", {t + 1: sales[t] for t in range(4)}),
       row("Variable costs", {t + 1: -vcs[t] for t in range(4)}),
       row("Incremental fixed costs", {t + 1: -fcs[t] for t in range(4)}),
       row("Taxable operating cash flow", {t + 1: op[t] for t in range(4)}),
       row("Tax at 25%", {t + 2: rows["taxop"][t] for t in range(4)}),
       row("Tax benefit of TAD", {t + 2: rows["taxtad"][t] for t in range(4)}),
       row("Machinery and scrap", {0: -COST, 4: SCRAP}),
       row("Working capital", {t: rows["wc"][t] for t in range(5)}),
       ["Net cash flow"] + [money(c) for c in cf],
       ["Discount factor at 11%", "1.000"] + [f"{df(.11, t):.3f}" for t in range(1, 6)],
       ["Present value"] + [money(p) for p in pv]]
w0 = 4.2
T(tbl, widths=[w0] + [(CONTENT_CM - w0) / 6] * 6, size=7.5, bold_rows=(4, 9, 11), right=range(1, 7), center=set())
P(f"**NPV = ${money(npv)}**, positive, so the project is financially acceptable.", after=8)

P("**Workings**", keep=True)
wk = [["Year", "1", "2", "3", "4"],
      ["Selling price ($48 × 1.03^n)"] + [f"{48 * 1.03 ** t:.2f}" for t in range(1, 5)],
      ["Variable cost ($27 × 1.05^n)"] + [f"{27 * 1.05 ** t:.2f}" for t in range(1, 5)],
      ["Fixed costs ($520,000 × 1.04^n)"] + [money(f) for f in fcs]]
T(wk, widths=[6.0] + [(CONTENT_CM - 6.0) / 4] * 4, size=8.5, right=range(1, 5), center=set())

tad_rows = [["Year", "TWDV b/f", "TAD / balancing adjustment", "Tax benefit at 25%", "Year received"]]
w = COST
for t in range(3):
    a = rows["tad"][t]
    tad_rows.append([str(t + 1), money(w), money(a), money(rows["taxtad"][t]), str(t + 2)])
    w -= a
tad_rows.append(["4", money(w), f"{money(rows['tad'][3])} (balancing allowance: {money(w)} − {money(SCRAP)})",
                 money(rows["taxtad"][3]), "5"])
T(tad_rows, widths=[1.3, 2.8, 6.4, 3.2, CONTENT_CM - 13.7], size=8.5, center={0, 4}, right={1, 3})
P(f"Check: total allowances = {money(COST)} − {money(SCRAP)} = {money(COST - SCRAP)}. No 25% writing-down allowance is claimed in Year 4; the balancing allowance replaces it.", size=9)

wc_rows = [["Year", "0", "1", "2", "3", "4"],
           ["Working capital held (10% of the coming year's sales)"] + [money(x) for x in rows["wcl"]] + ["nil"],
           ["Cash flow (increase)/release"] + [money(x) for x in rows["wc"]]]
T(wc_rows, widths=[5.6] + [(CONTENT_CM - 5.6) / 5] * 5, size=8.5, right=range(1, 6), center=set())

P("**Irrelevant items excluded:**", keep=True)
bullet("Market research of $75,000 is a **sunk cost**. It has already been paid, so the decision does not change it.")
bullet(f"Head office overheads of $90,000 are **not incremental**: they are incurred whether or not the project goes ahead. Including them (inflated at 4%) would reduce the NPV to ${money(npv_ovh)}, which is why the draft appraisal wrongly showed a negative result.")

P("**Marking guide (14 marks)**", keep=True, before=6)
T([["Item", "Marks"],
   ["Inflated selling price and sales revenue", "2"],
   ["Inflated variable costs", "1"],
   ["Inflated incremental fixed costs", "1"],
   ["Exclusion of sunk cost and allocated overheads (with reasons)", "1"],
   ["Tax on operating cash flows, one year in arrears", "1"],
   ["Tax-allowable depreciation, Years 1 to 3", "2"],
   ["Balancing allowance and its tax effect", "1"],
   ["Timing of TAD tax benefits (one year in arrears)", "1"],
   ["Working capital requirement and incremental flows", "2"],
   ["Machinery cost and scrap value", "1"],
   ["Discounting and NPV", "1"],
   ["**Total**", "**14**"]], widths=[CONTENT_CM - 2.5, 2.5], center={1})

P("**(b) Advice and risk assessment (6 marks)**", keep=True, before=6)
P("**Advice (up to 2 marks).** The revised NPV is positive at about $153,000, so the project is expected to increase shareholder wealth and should be accepted on financial grounds. The draft appraisal was wrong to include the $75,000 market research, which is a sunk cost, and the head office overhead allocation, which is not an incremental cash flow. Only future, incremental cash flows are relevant. The NPV is small relative to the $2.9m investment (about 5%), so the board should look carefully at how reliable the forecasts are before committing.")
P("**Sensitivity analysis (up to 2 marks).** This measures how far each key variable (selling price, volume, variable cost, fixed costs, cost of capital) could change before the NPV fell to zero. The sensitivity margin is NPV ÷ PV of the cash flows affected by that variable. Selling price is likely to be the most sensitive variable, because the PV of sales revenue is large compared with the NPV. That tells management where to concentrate forecasting and control effort. Its limitations are that it changes one variable at a time, ignores how likely each change is, and does not give a decision rule.")
P("**Probability analysis (up to 2 marks).** Probabilities could be assigned to different demand scenarios (for example low, expected and high volumes) to calculate an expected NPV and the probability of a negative NPV. That gives the board a measure of the risk of loss as well as the average expected return. However, the probabilities are subjective, and an expected value is a long-run average that is less meaningful for a one-off project of this size. Simulation could extend the analysis to allow several variables to change at the same time.")

H("Question 32: Calloway Hardware Co", 3, page_break=True)
open_block()
inv_d = 4.20 / 25.55 * 365
rec_d = 6.00 / 36.50 * 365
pay_d = 3.50 / 25.55 * 365
coc = inv_d + rec_d - pay_d
cur = (4.20 + 6.00) / (3.50 + 5.80)
quick = 6.00 / (3.50 + 5.80)
P("**(a) Working capital ratios (4 marks)**", keep=True)
T([["Measure", "Calculation", "Calloway", "Sector"],
   ["Inventory days", "4.20 / 25.55 × 365", f"{inv_d:.0f} days", ""],
   ["Receivables days", "6.00 / 36.50 × 365", f"{rec_d:.0f} days", ""],
   ["Payables days", "3.50 / 25.55 × 365", f"({pay_d:.0f}) days", ""],
   ["Cash operating cycle", "", f"{coc:.0f} days", "45 days"],
   ["Current ratio", "(4.20 + 6.00) / (3.50 + 5.80)", f"{cur:.2f} times", "1.50 times"],
   ["Quick ratio", "6.00 / (3.50 + 5.80)", f"{quick:.2f} times", "0.90 times"]],
  widths=[3.6, 6.4, 3.2, CONTENT_CM - 13.2], size=9, bold_rows=(4,), center={2, 3})
P("Calloway's cash operating cycle is 25 days longer than the sector's, mainly because of long inventory and receivables periods. Both liquidity ratios are well below the sector averages, which supports the bank's concern: current liabilities, including the overdraft, are only just covered by current assets.")

base_fin = 6.00e6 * 0.09
bd = 0.01 * 36.5e6
admin = 180_000
new_rec = 36.5e6 * 40 / 365
adv = 0.8 * new_rec
fac_int = adv * 0.10
od_int = (new_rec - adv) * 0.09
fee = 0.015 * 36.5e6
cur_cost = base_fin + bd + admin
new_cost = fac_int + od_int + fee
P("**(b) Evaluation of the factoring offer (8 marks)**", keep=True, before=6)
T([["", "Current ($)", "With factor ($)"],
   ["Receivables balance", money(6.0e6), f"{money(new_rec)} (36.5m × 40/365)"],
   ["Finance cost: overdraft at 9%", money(base_fin), f"{money(od_int)} (20% × {money(new_rec)} × 9%)"],
   ["Finance cost: factor advance at 10%", "nil", f"{money(fac_int)} (80% × {money(new_rec)} × 10%)"],
   ["Factor fee (1.5% × $36.5m)", "nil", money(fee)],
   ["Bad debts (1% × $36.5m)", money(bd), "nil (non-recourse)"],
   ["Credit control administration", money(admin), "nil (saved)"],
   ["**Total annual cost**", f"**{money(cur_cost)}**", f"**{money(new_cost)}**"]],
  widths=[6.2, 3.4, CONTENT_CM - 9.6], size=9, right={1}, center=set())
P(f"**Net annual benefit of factoring = {money(cur_cost)} − {money(new_cost)} = ${money(cur_cost - new_cost)}**, so the offer is financially acceptable.")
P("Liquidity would also improve: receivables would fall by $2.0m and the factor would advance $3.2m, so the overdraft could be cut by up to $5.2m, easing the bank's concern.")
P("**Common errors:**", keep=True)
bullet("Charging the factor's 10% on all receivables instead of the 80% advanced.")
bullet("Forgetting that the 20% not advanced still has to be financed by the overdraft.")
bullet("Treating bad debt savings as available under a *with-recourse* arrangement. Here the factoring is non-recourse, so the factor bears the bad debts.")
bullet("Calculating the fee on receivables instead of revenue.")
T([["Marking guide", "Marks"],
   ["Current financing cost of receivables", "1"], ["New receivables balance", "1"],
   ["Factor finance cost on the 80% advance", "1"], ["Overdraft cost on the remaining 20%", "1"],
   ["Factor fee", "1"], ["Bad debt saving", "1"], ["Administration saving", "1"],
   ["Net benefit and conclusion", "1"], ["**Total**", "**8**"]], widths=[CONTENT_CM - 2.5, 2.5], center={1})

P("**(c) Working capital financing policies (4 marks)**", keep=True, before=6)
P("Current assets are made up of **permanent** current assets (the minimum level always needed) and **fluctuating** current assets (seasonal or temporary peaks).")
bullet("**Matching policy:** long-term finance funds non-current assets and permanent current assets, and short-term finance funds fluctuating current assets. The maturity of the finance matches the life of the assets.")
bullet("**Conservative policy:** long-term finance funds all permanent assets and some of the fluctuating current assets. This is lower risk but more expensive, because long-term finance usually costs more and surplus funds may earn little.")
bullet("**Aggressive policy:** short-term finance funds all fluctuating current assets and some of the permanent current assets. This is cheaper, but it carries refinancing risk and interest rate risk.")
P("Calloway funds **all** of its current assets, including the permanent element, with its overdraft, which is repayable on demand. This is a highly **aggressive** policy. The bank's concern shows the risk: if the overdraft were withdrawn or reduced, Calloway would struggle to pay its suppliers.")

P("**(d) Non-financial factors before factoring (4 marks; 1 mark per well-explained point)**", keep=True, before=6)
for b in (
    "**Customer relationships:** the factor will contact Calloway's customers directly. An aggressive collection style could damage long-standing relationships, and customers may see the use of a factor as a sign of financial weakness.",
    "**Loss of control:** Calloway gives up control of its sales ledger and credit decisions. The factor may refuse to take on some customers, which could restrict sales.",
    "**Reputation of the factor:** Calloway should check the factor's standing, service levels and experience in the building supplies sector.",
    "**Exit and flexibility:** contracts often have minimum terms. Once the in-house credit control team has been made redundant, it is hard to bring the function back in-house.",
    "**Staff impact:** making credit control staff redundant may affect morale, and losing their knowledge of customers may be a disadvantage.",
    "**Alternatives:** Calloway could improve its own credit control, offer early settlement discounts, use invoice discounting (which keeps customer contact in-house) or arrange longer-term finance.",
):
    bullet(b)

P("**End of solution key**", align=WD_ALIGN_PARAGRAPH.CENTER, before=12)

finish(OUT_A, "ACCA FM Mock Exam 02: Solutions")
print("saved", OUT_Q, OUT_A, "NPV", round(npv), "NPV with overhead", round(npv_ovh), "factoring benefit", round(cur_cost - new_cost),
      "COC", round(coc), "ratios", round(cur, 2), round(quick, 2))
